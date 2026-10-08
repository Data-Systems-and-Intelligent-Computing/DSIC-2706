"""
update_dataset.py
1. Hapus file MP3 dari luar Indonesia (Malaysia)
2. Ambil Sampling Rate dari file audio (mutagen)
3. Ambil metadata lain dari Xeno-Canto API v3
4. Buat ulang Excel yang diperbarui
"""

import os, re, time, requests
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT_DIR    = r"D:\FILE AND TASK\TA\data\xeno_canto"
OUTPUT_FILE = r"D:\FILE AND TASK\TA\data\xeno_canto\Metadata_XenoCanto.xlsx"
API_KEY     = "95dc5b3f6834b56e17da4cf68cc890cb470d1284"

SPECIES_LIST = [
    "Apalharpactes mackloti",
    "Batrachostomus_poliolophus",
    "Caprimulgus_pulchellus",
    "Carpococcyx_viridis",
    "Erythropitta_venusta",
    "Gypsophila_rufipectus",
    "Hydrornis_schneideri",
    "Ixos_sumatranus",
    "Lophura_inornata",
    "Myophonus_melanurus",
    "Napothera_albostriata",
    "Niltava sumatrana",
    "Pellorneum_buettikoferi",
    "Picus_dedemi",
    "Polyplectron_chalcurum",
]

HEADERS_HTTP = {
    "X-API-Key" : API_KEY,
    "User-Agent": "BirdSoundResearch/1.0",
    "Accept"    : "application/json",
}

def log(msg):
    print(msg, flush=True)

def extract_xc_id(filename):
    match = re.match(r'XC(\d+)', filename, re.IGNORECASE)
    return match.group(1) if match else None

def get_sampling_rate(filepath):
    """Baca sampling rate dari file MP3 menggunakan mutagen (read-only)."""
    try:
        from mutagen.mp3 import MP3
        audio = MP3(filepath)
        return int(audio.info.sample_rate)
    except Exception as e:
        log(f"    [WARN] Gagal baca sampling rate: {e}")
        return ""

def get_xc_metadata_v3(xc_id, retries=3):
    urls = [
        f"https://xeno-canto.org/api/3/recordings?query=nr:{xc_id}&key={API_KEY}",
        f"https://xeno-canto.org/api/3/recordings/{xc_id}?key={API_KEY}",
    ]
    for attempt in range(retries):
        for url in urls:
            try:
                resp = requests.get(url, headers=HEADERS_HTTP, timeout=15)
                if resp.status_code == 200:
                    data = resp.json()
                    recs = data.get("recordings", [])
                    if not recs and "id" in data:
                        recs = [data]
                    if recs:
                        r = recs[0]
                        return {
                            "lat"      : r.get("lat", ""),
                            "lng"      : r.get("lng", ""),
                            "lokasi"   : r.get("loc", ""),
                            "negara"   : r.get("cnt", ""),
                            "kualitas" : r.get("q", ""),
                            "tipe"     : r.get("type", ""),
                            "tanggal"  : r.get("date", ""),
                            "rekorder" : r.get("rec", ""),
                        }
                elif resp.status_code == 401:
                    log(f"  [AUTH ERROR] XC{xc_id}: {resp.text[:100]}")
                    return None
            except Exception as e:
                log(f"  [Err percobaan {attempt+1}] XC{xc_id}: {e}")
                time.sleep(1)
        if attempt < retries - 1:
            time.sleep(2)
    return None

# ============================================================
# LANGKAH 1: Identifikasi & hapus file non-Indonesia
# ============================================================
log("=" * 60)
log("  LANGKAH 1: Identifikasi file luar Indonesia dari Excel")
log("=" * 60)

existing_excel = OUTPUT_FILE
files_to_delete = []

if os.path.exists(existing_excel):
    wb_old = openpyxl.load_workbook(existing_excel)
    ws_old = wb_old["SEMUA DATA"]
    headers_old = [cell.value for cell in ws_old[1]]

    for row in ws_old.iter_rows(min_row=2, values_only=True):
        data = dict(zip(headers_old, row))
        country = str(data.get("Negara", "")).strip()
        if country.lower() != "indonesia":
            species  = data.get("Spesies", "")
            subfolder= data.get("Sub-folder", "")
            fname    = data.get("Nama File", "")
            fpath    = os.path.join(ROOT_DIR, species, subfolder, fname)
            files_to_delete.append({
                "path"   : fpath,
                "negara" : country,
                "lokasi" : data.get("Lokasi", ""),
            })
else:
    log("  [WARN] Excel lama tidak ditemukan, akan scan ulang dari API.")

log(f"\n  File yang akan dihapus ({len(files_to_delete)} file):")
for f in files_to_delete:
    log(f"  - [{f['negara']}] {f['path']}")
    log(f"      Lokasi: {f['lokasi']}")

log("")
deleted_count = 0
for f in files_to_delete:
    if os.path.exists(f["path"]):
        os.remove(f["path"])
        log(f"  [HAPUS] {os.path.basename(f['path'])}")
        deleted_count += 1
    else:
        log(f"  [SKIP] File tidak ditemukan: {f['path']}")

log(f"\n  => {deleted_count} file berhasil dihapus dari folder.")

# ============================================================
# LANGKAH 2: Scan ulang folder & ambil data
# ============================================================
log(f"\n{'='*60}")
log("  LANGKAH 2: Scan folder & fetch metadata (dengan Sampling Rate)")
log("=" * 60)

# Install mutagen jika belum ada
try:
    from mutagen.mp3 import MP3
    log("  [OK] mutagen tersedia")
except ImportError:
    log("  [INFO] Menginstall mutagen...")
    import subprocess
    subprocess.run(["pip", "install", "mutagen", "-q"], check=True)
    from mutagen.mp3 import MP3
    log("  [OK] mutagen berhasil diinstall")

all_rows = []
species_files = {}

for species_folder in SPECIES_LIST:
    species_path = os.path.join(ROOT_DIR, species_folder)
    if not os.path.isdir(species_path):
        log(f"\n[SKIP] Folder tidak ditemukan: {species_folder}")
        continue

    log(f"\n{'-'*55}")
    log(f"[SPESIES] {species_folder}")
    rows_for_species = []

    for subfolder in sorted(os.listdir(species_path)):
        sub_path = os.path.join(species_path, subfolder)
        if not os.path.isdir(sub_path):
            continue

        for fname in sorted(os.listdir(sub_path)):
            if not fname.lower().endswith(".mp3"):
                continue

            xc_id    = extract_xc_id(fname)
            filepath = os.path.join(sub_path, fname)

            if not xc_id:
                log(f"  [WARN] Tidak bisa ekstrak XC ID: {fname}")
                continue

            short = fname[:50] + "..." if len(fname) > 50 else fname
            log(f"  [{subfolder}] XC{xc_id} -> {short}")

            # Sampling rate dari file audio langsung
            sr = get_sampling_rate(filepath)
            log(f"           Sampling Rate: {sr} Hz")

            # Metadata dari API
            meta = get_xc_metadata_v3(xc_id)

            row = {
                "Spesies"              : species_folder,
                "Sub-folder"           : subfolder,
                "Nama File"            : fname,
                "XC ID"                : f"XC{xc_id}",
                "Link XenoCanto"       : f"https://xeno-canto.org/{xc_id}",
                "Latitude"             : meta["lat"]      if meta else "GAGAL",
                "Longitude"            : meta["lng"]      if meta else "GAGAL",
                "Lokasi"               : meta["lokasi"]   if meta else "GAGAL",
                "Negara"               : meta["negara"]   if meta else "GAGAL",
                "Sampling Rate (Hz)"   : sr,
                "Kualitas (A-E)"       : meta["kualitas"] if meta else "GAGAL",
                "Tipe Suara"           : meta["tipe"]     if meta else "GAGAL",
                "Tanggal Rekaman"      : meta["tanggal"]  if meta else "GAGAL",
                "Rekorder"             : meta["rekorder"] if meta else "GAGAL",
            }
            rows_for_species.append(row)
            all_rows.append(row)
            time.sleep(0.4)

    species_files[species_folder] = rows_for_species
    log(f"  => {len(rows_for_species)} rekaman")

# ============================================================
# LANGKAH 3: Buat Excel baru
# ============================================================
log(f"\n{'='*60}")
log("  LANGKAH 3: Membuat file Excel...")

HEADER_COLORS = [
    "1F4E79","2E75B6","1C6B2B","375623","833C0B",
    "7B3F00","4A235A","154360","0E4D2E","5D2E0C",
    "1B2631","4A4A00","4A0000","00394A","2D0B4A",
]

COLUMNS    = [
    "Spesies","Sub-folder","Nama File","XC ID","Link XenoCanto",
    "Latitude","Longitude","Lokasi","Negara",
    "Sampling Rate (Hz)","Kualitas (A-E)","Tipe Suara",
    "Tanggal Rekaman","Rekorder",
]
COL_WIDTHS = [28,12,52,12,38,12,12,35,14,18,14,18,16,25]

thin   = Side(style="thin", color="CCCCCC")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(cell, hex_color):
    cell.font      = Font(bold=True, color="FFFFFF", size=11)
    cell.fill      = PatternFill("solid", fgColor=hex_color)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border    = border

def style_cell(cell, even=True):
    cell.fill      = PatternFill("solid", fgColor="F2F2F2" if even else "FFFFFF")
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    cell.border    = border

wb = openpyxl.Workbook()

# -- Sheet SEMUA DATA
ws_all = wb.active
ws_all.title = "SEMUA DATA"
ws_all.freeze_panes = "A2"
ws_all.row_dimensions[1].height = 30
for ci, (col, w) in enumerate(zip(COLUMNS, COL_WIDTHS), 1):
    c = ws_all.cell(row=1, column=ci, value=col)
    style_header(c, "1F4E79")
    ws_all.column_dimensions[get_column_letter(ci)].width = w

for ri, row in enumerate(all_rows, 2):
    even = (ri % 2 == 0)
    for ci, col in enumerate(COLUMNS, 1):
        val = row.get(col, "")
        c   = ws_all.cell(row=ri, column=ci, value=val)
        style_cell(c, even)
        if col == "Link XenoCanto" and isinstance(val, str) and val.startswith("http"):
            c.hyperlink = val
            c.font = Font(color="0563C1", underline="single")

# -- Sheet per Spesies
for si, (species, rows) in enumerate(species_files.items()):
    sheet_name = species.replace("_", " ")[:31]
    ws = wb.create_sheet(title=sheet_name)
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30
    hc = HEADER_COLORS[si % len(HEADER_COLORS)]
    for ci, (col, w) in enumerate(zip(COLUMNS, COL_WIDTHS), 1):
        c = ws.cell(row=1, column=ci, value=col)
        style_header(c, hc)
        ws.column_dimensions[get_column_letter(ci)].width = w
    if not rows:
        ws.cell(row=2, column=1, value="Tidak ada data")
        continue
    for ri, row in enumerate(rows, 2):
        even = (ri % 2 == 0)
        for ci, col in enumerate(COLUMNS, 1):
            val = row.get(col, "")
            c   = ws.cell(row=ri, column=ci, value=val)
            style_cell(c, even)
            if col == "Link XenoCanto" and isinstance(val, str) and val.startswith("http"):
                c.hyperlink = val
                c.font = Font(color="0563C1", underline="single")
    log(f"  Sheet '{sheet_name}' => {len(rows)} rekaman")

# -- Sheet RINGKASAN
ws_sum = wb.create_sheet(title="RINGKASAN", index=1)
ws_sum.freeze_panes = "A2"
sum_headers = ["No","Spesies","Jumlah Rekaman","Sub-folder","Negara (unik)","Sampling Rate Unik (Hz)"]
sum_widths  = [5, 32, 16, 22, 20, 30]
for ci, (col, w) in enumerate(zip(sum_headers, sum_widths), 1):
    c = ws_sum.cell(row=1, column=ci, value=col)
    style_header(c, "1F4E79")
    ws_sum.column_dimensions[get_column_letter(ci)].width = w
ws_sum.row_dimensions[1].height = 25

for si, (species, rows) in enumerate(species_files.items(), 1):
    subfolders = sorted(set(r["Sub-folder"] for r in rows))
    countries  = sorted(set(r["Negara"] for r in rows if r["Negara"] not in ("","GAGAL")))
    srates     = sorted(set(str(r["Sampling Rate (Hz)"]) for r in rows if r["Sampling Rate (Hz)"]))
    even = (si % 2 == 0)
    for ci, val in enumerate([
        si,
        species.replace("_"," "),
        len(rows),
        ", ".join(subfolders),
        ", ".join(countries) if countries else "-",
        ", ".join(srates) if srates else "-",
    ], 1):
        c = ws_sum.cell(row=si+1, column=ci, value=val)
        style_cell(c, even)

wb.save(OUTPUT_FILE)

log(f"\n{'='*60}")
log(f"  File Excel berhasil diperbarui:")
log(f"     {OUTPUT_FILE}")
log(f"  Total rekaman (setelah hapus luar negeri): {len(all_rows)}")
log(f"  File dihapus dari folder: {deleted_count}")
log("=" * 60)
