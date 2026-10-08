"""
check_countries.py - Cek dulu negara mana saja sebelum hapus
"""
import openpyxl

EXCEL_FILE = r"D:\FILE AND TASK\TA\data\xeno_canto\Metadata_XenoCanto.xlsx"

wb = openpyxl.load_workbook(EXCEL_FILE)
ws = wb["SEMUA DATA"]

headers = [cell.value for cell in ws[1]]
print("Headers:", headers)
print()

non_indo = []
for row in ws.iter_rows(min_row=2, values_only=True):
    data = dict(zip(headers, row))
    country = str(data.get("Negara", "")).strip()
    if country.lower() != "indonesia":
        non_indo.append({
            "Spesies"  : data.get("Spesies", ""),
            "Subfolder": data.get("Sub-folder", ""),
            "File"     : data.get("Nama File", ""),
            "Negara"   : country,
            "Lokasi"   : data.get("Lokasi", ""),
            "Lat"      : data.get("Latitude", ""),
            "Lng"      : data.get("Longitude", ""),
        })

print(f"Total rekaman BUKAN Indonesia: {len(non_indo)}")
print()
for r in non_indo:
    print(f"  [{r['Negara']}] {r['Spesies']} / {r['Subfolder']} / {r['File'][:60]}")
    print(f"         Lokasi: {r['Lokasi']}  | Lat: {r['Lat']} Lng: {r['Lng']}")
